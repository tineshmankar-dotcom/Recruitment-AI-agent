import io
import hashlib
from typing import Tuple
from pypdf import PdfReader
from docx import Document

class DocumentParser:
    @staticmethod
    def compute_hash(content: bytes) -> str:
        """Compute SHA-256 hash of file bytes for duplicate detection."""
        return hashlib.sha256(content).hexdigest()

    @staticmethod
    def extract_text(content: bytes, filename: str) -> Tuple[str, str]:
        """
        Extract text from file bytes based on file extension.
        Returns (extracted_text, file_type)
        """
        ext = filename.lower().split(".")[-1]
        
        if ext == "pdf":
            text = DocumentParser._extract_pdf(content)
            return text, "pdf"
        elif ext in ["docx", "doc"]:
            text = DocumentParser._extract_docx(content)
            return text, "docx"
        elif ext in ["txt", "text", "md"]:
            text = DocumentParser._extract_txt(content)
            return text, "txt"
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Supported: PDF, DOCX, TXT")

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        stream = io.BytesIO(content)
        reader = PdfReader(stream)
        text_parts = []
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        stream = io.BytesIO(content)
        doc = Document(stream)
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
                    
        return "\n".join(text_parts)

    @staticmethod
    def _extract_txt(content: bytes) -> str:
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")
